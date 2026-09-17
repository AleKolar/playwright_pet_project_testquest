import os
import re
from datetime import datetime

import openpyxl
import pytest
from playwright.sync_api import Page


# --------------------- Вспомогательная функция для извлечения числа ---------------------

def extract_price(text: str) -> float:
    """
    Извлекает число из строки с учётом обычного и типографского минуса.

    Примеры:
        '200 429 ₽' -> 200429.0
        '−10 021 ₽' -> -10021.0
        'включено' -> 0.0
    """
    match = re.search(
        r"([−\-]?[\d\s\u00A0]+(?:\.\d+)?)",
        text,
    )

    if match:
        num_str = match.group(1)
        num_str = re.sub(r"[\s\u00A0]+", "", num_str)
        num_str = num_str.replace("−", "-")
        return float(num_str)

    return 0.0


# --------------------- Настройка параметров ---------------------

def select_tariff(page: Page, tariff: str):
    """Выбрать тариф TestQuest Облако или TestQuest Коробка."""
    page.click(
        f"span.checkbox-text:has-text('{tariff}')"
    )


def set_licenses(page: Page, count: int):
    """Установить количество лицензий."""
    page.fill(
        "#licenceCount",
        str(count),
    )


def set_period(page: Page, months: int):
    """Установить срок в месяцах."""
    page.fill(
        "#termMonths",
        str(months),
    )


def set_discount(page: Page, percent: int):
    """Установить скидку в процентах."""
    page.fill(
        "#discountPercent",
        str(percent),
    )


def toggle_module(
    page: Page,
    module_name: str,
    enable: bool = True,
):
    """Включить или выключить модуль по названию."""
    li = page.locator(
        f"li.module-item:has("
        f"label.module-label:has-text('{module_name}')"
        f")"
    )

    checkbox = li.locator(
        "input[type='checkbox']:not([id*='contract'])"
    )

    if enable:
        checkbox.check()
    else:
        checkbox.uncheck()

# --------------------- Работа с датами ---------------------
def set_kp_issue_date(page: Page, date: str):
    """Установить дату выставления КП. Формат: YYYY-MM-DD."""
    page.fill("#kpIssueDate", date)


def set_kp_valid_until(page: Page, date: str):
    """Установить дату окончания действия КП. Формат: YYYY-MM-DD."""
    page.fill("#kpValidUntil", date)


def get_kp_issue_date(page: Page) -> str:
    """Получить дату выставления КП. Формат: YYYY-MM-DD."""
    return page.locator("#kpIssueDate").input_value()


def get_kp_valid_until(page: Page) -> str:
    """Получить дату окончания действия КП. Формат: YYYY-MM-DD."""
    return page.locator("#kpValidUntil").input_value()

# --------------------- Включить или выключить модуль по id чек-бокса ---------------------
def toggle_module_by_id(
    page: Page,
    module_id: str,
    enable: bool = True,
):
    """
    Включить или выключить модуль по id чек-бокса.

    Пример:
        module_id = "module-process-builder"
    """
    checkbox = page.locator(
        f"input[type='checkbox']#{module_id}"
    )

    if enable:
        checkbox.check()
    else:
        checkbox.uncheck()


# --------------------- Вкладки ---------------------

def switch_to_tab_implementation(page: Page):
    """Переключиться на вкладку «Внедрение»."""
    page.click("#config-tab-implementation")
    page.wait_for_timeout(300)


def switch_to_tab_tm(page: Page):
    """Переключиться на вкладку T&M."""
    page.click("#config-tab-tm")
    page.wait_for_timeout(300)


# --------------------- Внедрение ---------------------

def toggle_implementation_work(
    page: Page,
    work_name: str,
    enable: bool = True,
):
    """Включить или выключить работу внедрения."""
    switch_to_tab_implementation(page)

    li = page.locator(
        f"li.implementation-item:has("
        f"label:has-text('{work_name}')"
        f")"
    )

    if li.count() == 0:
        li = page.locator(
            f"li:has(label:has-text('{work_name}'))"
        ).first

    checkbox = li.locator(
        "input[type='checkbox']"
    ).first

    if enable:
        checkbox.check()
    else:
        checkbox.uncheck()

    page.wait_for_timeout(200)


def get_selected_implementation_works(
    page: Page,
) -> list:
    """Получить список выбранных работ внедрения."""
    switch_to_tab_implementation(page)

    items = page.locator(
        ".selected-work-name, .chosen-work-name"
    ).all()

    return [
        item.inner_text().strip()
        for item in items
    ]


# --------------------- T&M ---------------------

def get_tandm_hours(page: Page) -> int:
    """Получить количество часов T&M."""
    switch_to_tab_tm(page)

    field = page.locator(
        "#tmHours, #tm-hours, input[name='tmHours']"
    ).first

    if field.count() == 0:
        return 0

    return int(field.input_value() or 0)


def set_tandm_hours(
    page: Page,
    hours: int,
):
    """Установить количество часов T&M."""
    switch_to_tab_tm(page)

    field = page.locator(
        "#tmHours, #tm-hours, input[name='tmHours']"
    ).first

    if field.count() == 0:
        pytest.fail(
            "Поле часов T&M не найдено."
        )

    field.fill(str(hours))
    page.wait_for_timeout(300)


def get_tandm_tariff(
    page: Page,
) -> float:
    """
    Прочитать тариф T&M из текста вкладки.

    Ожидаемый текст:
        Тариф 5 850 ₽ / чч без НДС
    """
    switch_to_tab_tm(page)

    body_text = page.locator("body").inner_text()

    match = re.search(
        r"Тариф\s+([\d\s\u00A0]+)\s*₽\s*/\s*чч",
        body_text,
        re.IGNORECASE,
    )

    if not match:
        return 0.0

    return extract_price(
        match.group(1)
    )


# --------------------- Ожидание расчёта ---------------------

def wait_for_calculation(page: Page):
    """Дождаться появления итоговой суммы."""
    page.wait_for_selector(
        ".selected-summary-value",
        state="visible",
        timeout=10000,
    )

    page.wait_for_timeout(500)


def get_total_sum(
    page: Page,
) -> float:
    """Получить первую итоговую сумму из UI."""
    element = page.locator(
        ".selected-summary-value"
    ).first

    return extract_price(
        element.inner_text()
    )


# --------------------- Выгрузка Excel ---------------------

def download_excel(
    page: Page,
) -> str:
    """
    Нажать реальную кнопку «Выгрузить смету в Excel»,
    скачать Excel и сохранить его в корень проекта
    под уникальным именем.
    """
    project_root = os.path.dirname(
        os.path.dirname(
            os.path.abspath(__file__)
        )
    )

    with page.expect_download() as download_info:
        page.get_by_role(
            "button",
            name="Выгрузить смету в Excel",
        ).click()

    download = download_info.value

    original_name = (
        download.suggested_filename
        or "estimate.xlsx"
    )

    base_name, extension = os.path.splitext(
        original_name
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S_%f"
    )

    filename = (
        f"{base_name}_{timestamp}{extension}"
    )

    file_path = os.path.join(
        project_root,
        filename,
    )

    download.save_as(file_path)

    return file_path


# --------------------- Excel parsing ---------------------

def _extract_number_from_cell(
    cell,
) -> float | None:
    """Извлечь число из Excel-ячейки."""
    if cell.value is None:
        return None

    if isinstance(
        cell.value,
        (int, float),
    ):
        return float(cell.value)

    if isinstance(
        cell.value,
        str,
    ):
        cleaned = re.sub(
            r"[^\d\-.,]",
            "",
            cell.value,
        )

        cleaned = cleaned.replace(
            ",",
            ".",
        )

        try:
            return float(cleaned)
        except ValueError:
            return None

    return None


def get_excel_rows(
    file_path: str,
) -> list:
    """Получить все непустые строки активного листа."""
    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
    )

    sheet = workbook.active
    rows = []

    for row_number, row in enumerate(
        sheet.iter_rows(values_only=True),
        start=1,
    ):
        if any(
            cell is not None
            and str(cell).strip() != ""
            for cell in row
        ):
            rows.append(
                {
                    "row": row_number,
                    "cells": list(row),
                }
            )

    return rows


def find_excel_rows(
    file_path: str,
    search_text: str,
) -> list:
    """Найти строки Excel по тексту."""
    needle = search_text.lower()
    result = []

    for row in get_excel_rows(file_path):
        if any(
            cell is not None
            and needle in str(cell).lower()
            for cell in row["cells"]
        ):
            result.append(row)

    return result


def count_excel_rows(
    file_path: str,
    search_text: str,
) -> int:
    """Посчитать строки с указанным текстом."""
    return len(
        find_excel_rows(
            file_path,
            search_text,
        )
    )


def get_excel_total_rows(
    file_path: str,
) -> list:
    """Получить строки, содержащие «Всего» или «Итого»."""
    return (
        find_excel_rows(
            file_path,
            "всего",
        )
        + find_excel_rows(
            file_path,
            "итого",
        )
    )


def get_excel_row_amounts(
    file_path: str,
    label: str,
) -> dict | None:
    """
    Получить суммы из строки Excel.

    C = без НДС
    D = НДС
    E = с НДС
    """
    rows = find_excel_rows(
        file_path,
        label,
    )

    if not rows:
        return None

    cells = rows[0]["cells"]

    if len(cells) < 5:
        return None

    class Cell:
        def __init__(
            self,
            value,
        ):
            self.value = value

    return {
        "without_vat": _extract_number_from_cell(
            Cell(cells[2])
        ),
        "vat": _extract_number_from_cell(
            Cell(cells[3])
        ),
        "with_vat": _extract_number_from_cell(
            Cell(cells[4])
        ),
    }


def get_all_excel_text(
    file_path: str,
) -> str:
    """Получить весь текст Excel одной строкой."""
    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
    )

    sheet = workbook.active
    parts = []

    for row in sheet.iter_rows(
        values_only=True,
    ):
        for cell in row:
            if cell is not None:
                parts.append(str(cell))

    return " | ".join(parts)

# --------------------- Разбивка по позициям ---------------------

def get_breakdown(page: Page) -> dict:
    """
    Получить разбивку стоимости выбранных позиций.

    Итоговые строки «Всего» и «Лицензии со скидкой» не учитываются.
    T&M, резерв, скидка и другие служебные строки также исключаются.
    """
    breakdown = {}

    exclude_keywords = [
        "Резерв",
        "T&M",
        "Time",
        "Скидка",
    ]

    # Позиции.
    name_elements = page.locator(
        ".selected-item-name"
    ).all()

    for name_element in name_elements:
        parent = name_element.locator("xpath=..")
        price_element = parent.locator(
            ".selected-item-price"
        )

        if price_element.count() == 0:
            continue

        price_text = price_element.inner_text().strip()

        if "включено" in price_text.lower():
            continue

        price = extract_price(price_text)

        if price == 0:
            continue

        name = name_element.inner_text().strip()

        if any(
            keyword in name
            for keyword in exclude_keywords
        ):
            continue

        breakdown[name] = price

    # Итоговые строки.
    label_elements = page.locator(
        ".selected-summary-label"
    ).all()

    for label_element in label_elements:
        parent = label_element.locator("xpath=..")
        value_element = parent.locator(
            ".selected-summary-value"
        )

        if value_element.count() == 0:
            continue

        value_text = value_element.inner_text().strip()
        price = extract_price(value_text)

        if price == 0:
            continue

        name = label_element.inner_text().strip()

        if name in (
            "Всего",
            "Лицензии со скидкой",
        ):
            continue

        if any(
            keyword in name
            for keyword in exclude_keywords
        ):
            continue

        breakdown[name] = price

    return breakdown

# --------------------- Проверка Excel ---------------------

def validate_excel(
    file_path: str,
    expected_total: float,
) -> bool:
    """
    Проверить, что Excel содержит ожидаемую итоговую сумму.
    """
    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
    )

    sheet = workbook.active

    # Сначала ищем строки «Итого» / «Всего».
    for row in sheet.iter_rows(
        min_row=1,
        max_row=sheet.max_row,
        max_col=sheet.max_column,
    ):
        for cell in row:
            if not isinstance(cell.value, str):
                continue

            cell_text = cell.value.lower()

            if (
                "итого" not in cell_text
                and "всего" not in cell_text
            ):
                continue

            # Проверяем саму ячейку.
            number = _extract_number_from_cell(cell)

            if (
                number is not None
                and abs(number - expected_total) < 0.01
            ):
                return True

            # Проверяем остальные ячейки строки.
            for other_cell in row:
                if other_cell == cell:
                    continue

                number = _extract_number_from_cell(
                    other_cell
                )

                if (
                    number is not None
                    and abs(number - expected_total) < 0.01
                ):
                    return True

    # Если итоговая строка не найдена,
    # ищем ожидаемую сумму во всём Excel.
    for row in sheet.iter_rows(
        min_row=1,
        max_row=sheet.max_row,
        max_col=sheet.max_column,
    ):
        for cell in row:
            number = _extract_number_from_cell(cell)

            if (
                number is not None
                and abs(number - expected_total) < 0.01
            ):
                return True

    return False

